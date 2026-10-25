"""
Professional Word Document (.docx) Generator for RoadSense Cloud Implementation Report.
Generates: RoadSense_Cloud_Architecture_Report.docx
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text, title="KEY ARCHITECTURAL HIGHLIGHT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F9FF")
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    # Left border bar in azure blue
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="0284C7"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"[{title}]  ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(2, 132, 199)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph()

def add_code_block(doc, code_str, filename=""):
    if filename:
        p_name = doc.add_paragraph()
        p_name.paragraph_format.space_before = Pt(6)
        p_name.paragraph_format.space_after = Pt(2)
        r_name = p_name.add_run(f"Source Listing: {filename}")
        r_name.bold = True
        r_name.font.name = "Consolas"
        r_name.font.size = Pt(9)
        r_name.font.color.rgb = RGBColor(100, 116, 139)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="CBD5E1"/>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

def style_heading_1(doc, text):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(18)
    r.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy
    return h

def style_heading_2(doc, text):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(13.5)
    r.bold = True
    r.font.color.rgb = RGBColor(109, 40, 217) # Royal Purple
    return h

def style_heading_3(doc, text):
    h = doc.add_heading(level=3)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_body_paragraph(doc, text, bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10.5)
        r_b.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10)
        r_b.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_row = tbl.rows[0]
    for idx, header in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "1E3A8A") # Navy
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 41, 59)
            
    # Apply column widths
    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    doc.add_paragraph()

def add_image_figure(doc, img_path, caption, width_in=6.2):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption)
        r_cap.italic = True
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(8.5)
def add_evidence_placeholder(doc, resource_title, resource_arn, details_dict, instruction=""):
    """
    Renders an official, beautifully styled Evidence Container for AWS Resources
    with dashed border, metadata tags, and explicit screenshot placement instructions.
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=180, bottom=180, left=240, right=240)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:top w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:right w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    
    r_hdr = p.add_run(f"📷 EVIDENCE SCREENSHOT PLACEHOLDER: {resource_title.upper()}\n")
    r_hdr.bold = True
    r_hdr.font.name = "Calibri"
    r_hdr.font.size = Pt(11)
    r_hdr.font.color.rgb = RGBColor(37, 99, 235)
    
    p_meta = cell.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_meta.paragraph_format.space_before = Pt(2)
    p_meta.paragraph_format.space_after = Pt(2)
    
    r_arn_lbl = p_meta.add_run("Target Resource Identifier / ARN: ")
    r_arn_lbl.bold = True
    r_arn_lbl.font.size = Pt(9.5)
    r_arn = p_meta.add_run(f"{resource_arn}\n")
    r_arn.font.name = "Consolas"
    r_arn.font.size = Pt(9)
    r_arn.font.color.rgb = RGBColor(71, 85, 105)
    
    for k, v in details_dict.items():
        rk = p_meta.add_run(f"• {k}: ")
        rk.bold = True
        rk.font.size = Pt(9)
        rv = p_meta.add_run(f"{v}\n")
        rv.font.size = Pt(9)
        rv.font.color.rgb = RGBColor(71, 85, 105)
        
    p_instr = cell.add_paragraph()
    p_instr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_instr.paragraph_format.space_before = Pt(14)
    p_instr.paragraph_format.space_after = Pt(10)
    
    r_box = p_instr.add_run("┌────────────────────────────────────────────────────────────────────────────────────────┐\n")
    r_box.font.name = "Consolas"
    r_box.font.size = Pt(8.5)
    r_box.font.color.rgb = RGBColor(148, 163, 184)
    
    r_in = p_instr.add_run("PASTE YOUR AWS CONSOLE SCREENSHOT HERE\n")
    r_in.bold = True
    r_in.font.name = "Calibri"
    r_in.font.size = Pt(11)
    r_in.font.color.rgb = RGBColor(71, 85, 105)
    
    r_desc = p_instr.add_run(instruction or "Paste screenshot showing active resource, configurations, and metric graphs.\n")
    r_desc.font.italic = True
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = RGBColor(100, 116, 139)
    
    r_box2 = p_instr.add_run("└────────────────────────────────────────────────────────────────────────────────────────┘")
    r_box2.font.name = "Consolas"
    r_box2.font.size = Pt(8.5)
    r_box2.font.color.rgb = RGBColor(148, 163, 184)
    
    doc.add_paragraph()

print("Docx generator helper functions initialized successfully.")
